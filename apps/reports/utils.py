import io
import math
import random
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from django.db import transaction
from django.db.models import Count, Q
from django.http import HttpResponse
from django.utils import timezone
from apps.master_data.models import Student
from apps.exam_engine.models import ExamSession, StudentAnswer
from apps.reports.models import StudentLogistic


FORMULA_PREFIXES = ('=', '+', '-', '@', '\t', '\r')


def excel_safe(value):
    """Teks yang diawali karakter formula ditulis sebagai teks biasa, bukan rumus Excel."""
    if isinstance(value, str) and value.startswith(FORMULA_PREFIXES):
        return "'" + value
    return value


def _now_wib():
    return timezone.localtime().strftime('%d-%m-%Y %H:%M')


def get_exam_related_schedules(exam):
    """Mengambil semua schedule ID yang merujuk ke paket ujian ini."""
    return list(exam.schedules.values_list('id', flat=True))


def get_student_best_session_map(exam, schedule_ids=None):
    """
    Mengambil map sesi terbaik per siswa untuk suatu paket ujian.
    Prioritas: finished > locked > active > started_at terbaru.
    """
    if schedule_ids is None:
        schedule_ids = get_exam_related_schedules(exam)

    if not schedule_ids:
        return {}

    # Hitungan benar/terjawab/total dianotasi sekaligus agar rekap tidak menjalankan query per siswa
    sessions = ExamSession.objects.filter(
        schedule_id__in=schedule_ids
    ).select_related('student', 'schedule').annotate(
        correct_n=Count('student_answers', filter=Q(student_answers__is_correct=True)),
        answered_n=Count('student_answers', filter=Q(student_answers__chosen_answer__isnull=False) & ~Q(student_answers__chosen_answer='')),
        total_n=Count('student_answers'),
    ).order_by('-started_at')

    session_map = {}
    for s in sessions:
        student_id = str(s.student_id)
        if student_id not in session_map:
            session_map[student_id] = s
        else:
            current = session_map[student_id]
            # finished prioritizes over non-finished
            if s.status == 'finished' and current.status != 'finished':
                session_map[student_id] = s
            # locked prioritizes over active if not finished
            elif s.status == 'locked' and current.status == 'active':
                session_map[student_id] = s

    return session_map


def _discrimination_label(d):
    if d >= 0.40:
        return "Sangat baik"
    if d >= 0.30:
        return "Baik"
    if d >= 0.20:
        return "Cukup"
    return "Perlu revisi"


def calculate_distractor_analysis(exam, schedule_ids=None, student_ids=None):
    """
    Analisis butir: tingkat kesukaran, daya pembeda (kelompok atas vs bawah 27%), dan sebaran pilihan A-E.
    Jawaban tersimpan dalam huruf asli soal, jadi sebaran tetap valid walau opsi diacak per siswa.
    student_ids: batasi ke siswa tertentu (guru hanya melihat kelas yang ia ajar).
    """
    if schedule_ids is None:
        schedule_ids = get_exam_related_schedules(exam)

    # 1. Ambil daftar butir soal ujian berurutan
    exam_questions = exam.exam_questions.select_related('question').order_by('order_number')
    questions = [eq.question for eq in exam_questions if eq.question]

    # 2. Ambil semua sesi yang selesai
    finished_sessions = ExamSession.objects.filter(schedule_id__in=schedule_ids, status='finished')
    if student_ids is not None:
        finished_sessions = finished_sessions.filter(student_id__in=student_ids)

    ranked = list(finished_sessions.order_by('-score').values_list('id', flat=True))
    finished_session_ids = ranked
    total_finished = len(ranked)
    group_size = max(1, round(total_finished * 0.27)) if total_finished else 0
    upper_ids = {str(i) for i in ranked[:group_size]}
    lower_ids = {str(i) for i in ranked[-group_size:]} if group_size else set()

    # 3. Ambil seluruh jawaban siswa pada sesi selesai
    answers = []
    if finished_session_ids:
        answers = list(
            StudentAnswer.objects.filter(
                session_id__in=finished_session_ids
            ).values('session_id', 'question_id', 'chosen_answer', 'is_correct')
        )

    # Group jawaban by question_id
    answers_by_question = {}
    for a in answers:
        qid = str(a['question_id'])
        if qid not in answers_by_question:
            answers_by_question[qid] = []
        answers_by_question[qid].append(a)

    analysis_results = []
    for idx, q in enumerate(questions, start=1):
        qid = str(q.id)
        q_answers = answers_by_question.get(qid, [])
        
        correct_count = 0
        wrong_count = 0
        blank_count = 0
        distro = {'A': 0, 'B': 0, 'C': 0, 'D': 0, 'E': 0}
        
        # Track answers per session to avoid duplicate counts
        session_choice_map = {}
        for ans in q_answers:
            sid = str(ans['session_id'])
            if sid not in session_choice_map:
                session_choice_map[sid] = ans['chosen_answer']

        kunci = (q.correct_answer or '').strip().upper()

        for sid in finished_session_ids:
            sid_str = str(sid)
            chosen = session_choice_map.get(sid_str)
            if not chosen:
                blank_count += 1
            else:
                chosen_clean = str(chosen).strip().upper()
                if chosen_clean in distro:
                    distro[chosen_clean] += 1
                
                if chosen_clean == kunci:
                    correct_count += 1
                else:
                    wrong_count += 1

        # Ketuntasan / Tingkat Kesukaran
        if total_finished > 0:
            percentage = round((correct_count / total_finished) * 100, 1)
            p_index = correct_count / total_finished
        else:
            percentage = 0.0
            p_index = 0.0

        if p_index >= 0.70:
            difficulty_level = "Mudah"
            difficulty_badge = "bg-emerald-100 text-emerald-700 dark:bg-emerald-950/40 dark:text-emerald-400"
        elif p_index >= 0.30:
            difficulty_level = "Sedang"
            difficulty_badge = "bg-blue-100 text-blue-700 dark:bg-blue-950/40 dark:text-blue-400"
        else:
            difficulty_level = "Sukar"
            difficulty_badge = "bg-red-100 text-red-700 dark:bg-red-950/40 dark:text-red-400"

        upper_correct = sum(1 for sid in upper_ids if str(session_choice_map.get(sid) or '').strip().upper() == kunci)
        lower_correct = sum(1 for sid in lower_ids if str(session_choice_map.get(sid) or '').strip().upper() == kunci)
        discrimination = round((upper_correct - lower_correct) / group_size, 2) if group_size else 0.0

        distractor_summary = []
        for opt in ['A', 'B', 'C', 'D', 'E']:
            if opt != kunci:
                cnt = distro[opt]
                distractor_summary.append(f"{opt}: {cnt}")

        analysis_results.append({
            'no': idx,
            'question_id': str(q.id),
            'question_text': q.question_text,
            'correct_answer': kunci,
            'correct': correct_count,
            'wrong': wrong_count,
            'blank': blank_count,
            'distro': distro,
            'correct_percentage': percentage,
            'difficulty_level': difficulty_level,
            'difficulty_badge': difficulty_badge,
            'p_index': round(p_index, 2),
            'discrimination': discrimination,
            'discrimination_label': _discrimination_label(discrimination) if total_finished else '-',
            'distractor_summary': ', '.join(distractor_summary),
        })

    return {
        'total_finished': total_finished,
        'questions_count': len(questions),
        'analysis_data': analysis_results
    }


class LogisticsCapacityError(ValueError):
    pass


def auto_distribute_logistics(levels, room_count, capacity, sessions_count, exam_period=''):
    """
    Mengocok dan mengalokasikan seluruh siswa aktif pada jenjang terpilih ke ruang dan sesi.
    Ditolak bila kapasitas (ruang x kursi x sesi) kurang dari jumlah siswa, supaya tidak ada siswa
    yang diam-diam kehilangan tempat. Hapus dan buat ulang berjalan dalam satu transaksi.
    """
    if not levels or room_count <= 0 or capacity <= 0 or sessions_count <= 0:
        return 0

    # Tarik seluruh siswa aktif di jenjang terkait
    students_qs = Student.objects.filter(
        status='aktif',
        class_room__level__in=levels
    ).select_related('class_room').order_by('full_name')

    students_list = list(students_qs)
    if not students_list:
        return 0

    total_capacity = room_count * capacity * sessions_count
    if total_capacity < len(students_list):
        raise LogisticsCapacityError(
            f"Kapasitas {room_count} ruang x {capacity} kursi x {sessions_count} sesi = {total_capacity} kursi, "
            f"kurang dari {len(students_list)} siswa. Tambah ruang, kursi, atau sesi."
        )

    random.shuffle(students_list)
    student_ids = [s.id for s in students_list]

    new_logistics = []
    student_idx = 0
    total_students = len(students_list)

    for s in range(1, sessions_count + 1):
        for r in range(1, room_count + 1):
            for c in range(1, capacity + 1):
                if student_idx < total_students:
                    st = students_list[student_idx]
                    new_logistics.append(
                        StudentLogistic(
                            student=st,
                            room_name=f"RUANG {r:02d}",
                            session_name=f"SESI {s}",
                            exam_period=exam_period
                        )
                    )
                    student_idx += 1
                else:
                    break

    with transaction.atomic():
        StudentLogistic.objects.filter(student_id__in=student_ids).delete()
        StudentLogistic.objects.bulk_create(new_logistics)

    return len(new_logistics)


def generate_exam_results_excel(exam, participants_data, analysis_data, kkm=75):
    """
    Membuat file Excel rekap hasil ujian lengkap dengan Rekap Nilai & Analisis Butir Soal.
    """
    wb = Workbook()
    
    # Styles
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    bold_font = Font(name="Arial", size=10, bold=True)
    regular_font = Font(name="Arial", size=10)
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    center_align = Alignment(horizontal='center', vertical='center')
    left_align = Alignment(horizontal='left', vertical='center')
    right_align = Alignment(horizontal='right', vertical='center')

    # Sheet 1: Rekap Nilai
    ws_rekap = wb.active
    ws_rekap.title = "Rekap Nilai"

    # Title Banner
    ws_rekap.append([f"REKAPITULASI HASIL UJIAN: {exam.title.upper()}"])
    ws_rekap.append([f"Mata Pelajaran: {exam.subject.name if exam.subject else '-'} | Tipe: {exam.exam_type} | Waktu Unduh: {_now_wib()} WIB"])
    ws_rekap.append([])  # empty row

    rekap_headers = [
        "No", "NIS", "Nama Lengkap Siswa", "Kelas", "Status Ujian", 
        "Jml Benar", "Jml Salah", "Kosong", "Nilai Akhir", f"Status Kelulusan (KKM {kkm:g})"
    ]
    ws_rekap.append(rekap_headers)

    # Style Header
    header_row_idx = 4
    for col_idx in range(1, len(rekap_headers) + 1):
        cell = ws_rekap.cell(row=header_row_idx, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border

    # Append Data
    for idx, p in enumerate(participants_data, start=1):
        row_data = [
            idx,
            p.get('nis', ''),
            p.get('full_name', '').upper(),
            p.get('class_name', '-'),
            p.get('status_label', 'Belum Mulai'),
            p.get('correct_count', 0),
            p.get('wrong_count', 0),
            p.get('blank_count', 0),
            p.get('score', 0),
            "TUNTAS" if p.get('score', 0) >= kkm and p.get('is_finished') else ("BELUM TUNTAS" if p.get('is_finished') else "-")
        ]
        ws_rekap.append([excel_safe(v) for v in row_data])
        current_row = header_row_idx + idx
        for col_idx in range(1, len(row_data) + 1):
            cell = ws_rekap.cell(row=current_row, column=col_idx)
            cell.font = regular_font
            cell.border = thin_border
            if col_idx in [1, 2, 4, 5, 6, 7, 8, 10]:
                cell.alignment = center_align
            elif col_idx == 3:
                cell.alignment = left_align
            elif col_idx == 9:
                cell.alignment = right_align
                cell.font = bold_font

    # Column Width Auto-Fit
    for col in ws_rekap.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws_rekap.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # Sheet 2: Analisis Butir Soal
    ws_analisis = wb.create_sheet(title="Analisis Pengecoh")
    ws_analisis.append([f"ANALISIS BUTIR SOAL & SEBARAN PENGECOH: {exam.title.upper()}"])
    ws_analisis.append([f"Mata Pelajaran: {exam.subject.name if exam.subject else '-'} | Jumlah Soal: {len(analysis_data)} Butir"])
    ws_analisis.append([])

    analisis_headers = [
        "No Soal", "Kunci", "Cuplikan Soal", "Benar", "Salah", "Kosong",
        "Pilih A", "Pilih B", "Pilih C", "Pilih D", "Pilih E",
        "Ketuntasan (%)", "Tingkat Kesukaran", "Daya Pembeda", "Kategori Daya Pembeda"
    ]
    ws_analisis.append(analisis_headers)

    header_row_idx = 4
    for col_idx in range(1, len(analisis_headers) + 1):
        cell = ws_analisis.cell(row=header_row_idx, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border

    for idx, item in enumerate(analysis_data, start=1):
        distro = item.get('distro', {})
        row_data = [
            item.get('no', idx),
            item.get('correct_answer', 'A'),
            (item.get('question_text') or '')[:80] + ('...' if len(item.get('question_text') or '') > 80 else ''),
            item.get('correct', 0),
            item.get('wrong', 0),
            item.get('blank', 0),
            distro.get('A', 0),
            distro.get('B', 0),
            distro.get('C', 0),
            distro.get('D', 0),
            distro.get('E', 0),
            f"{item.get('correct_percentage', 0)}%",
            item.get('difficulty_level', '-'),
            item.get('discrimination', 0),
            item.get('discrimination_label', '-'),
        ]
        ws_analisis.append([excel_safe(v) for v in row_data])
        current_row = header_row_idx + idx
        for col_idx in range(1, len(row_data) + 1):
            cell = ws_analisis.cell(row=current_row, column=col_idx)
            cell.font = regular_font
            cell.border = thin_border
            if col_idx != 3:
                cell.alignment = center_align
            else:
                cell.alignment = left_align

    for col in ws_analisis.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws_analisis.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # Return HTTP Response
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"Hasil_Ujian_{exam.subject.name if exam.subject else 'Mapel'}_{exam.title}.xlsx".replace(' ', '_').replace('/', '-')
    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


def generate_logistics_excel(logistics_qs):
    """
    Membuat file Excel rekap alokasi logistik ruangan dan sesi siswa.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Logistik Ruang & Sesi"

    header_fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    regular_font = Font(name="Arial", size=10)
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    ws.append(["DAFTAR ALOKASI LOGISTIK RUANGAN & SESI UJIAN"])
    ws.append([f"Waktu Unduh: {_now_wib()} WIB"])
    ws.append([])

    headers = ["No", "NIS", "Nama Lengkap Siswa", "Kelas", "Ruangan", "Sesi"]
    ws.append(headers)

    header_row_idx = 4
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=header_row_idx, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border

    for idx, log in enumerate(logistics_qs, start=1):
        row_data = [
            idx,
            log.student.nis if log.student else '',
            log.student.full_name.upper() if log.student else '',
            log.student.class_room.name if log.student and log.student.class_room else '',
            log.room_name,
            log.session_name,
        ]
        ws.append([excel_safe(v) for v in row_data])
        current_row = header_row_idx + idx
        for col_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.font = regular_font
            cell.border = thin_border
            if col_idx in [1, 2, 4, 5, 6]:
                cell.alignment = Alignment(horizontal='center', vertical='center')
            else:
                cell.alignment = Alignment(horizontal='left', vertical='center')

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"Alokasi_Logistik_{timezone.localtime().strftime('%Y%m%d_%H%M')}.xlsx"
    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response
