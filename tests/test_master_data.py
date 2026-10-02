"""Tests for Master Data Module & OpenAPI Synchronization (Task 4)."""
import io
import json
import uuid
from unittest.mock import MagicMock, patch
import pytest
import openpyxl
from django.contrib.auth import get_user_model
from django.test import Client, RequestFactory
from django.urls import reverse

from apps.accounts.models import Role
from apps.master_data.models import Major, ClassRoom, Teacher, Student, Subject, TeacherAssignment
from apps.master_data.services.openapi_sync import MasterDataSyncService, master_data_sync_service

User = get_user_model()


# ==============================================================================
# 1. Models & CRUD Operations Tests
# ==============================================================================

@pytest.mark.django_db
def test_major_model_crud():
    """Verify Major model CRUD, string representation, and uniqueness."""
    major = Major.objects.create(
        code='RPL',
        name='Rekayasa Perangkat Lunak'
    )
    assert str(major) == 'RPL - Rekayasa Perangkat Lunak'
    assert major.id is not None

    # Update
    major.name = 'Pengembangan Perangkat Lunak dan Gim'
    major.save()
    updated = Major.objects.get(code='RPL')
    assert updated.name == 'Pengembangan Perangkat Lunak dan Gim'

    # Uniqueness constraint
    with pytest.raises(Exception):
        Major.objects.create(code='RPL', name='Duplicate RPL')


@pytest.mark.django_db
def test_classroom_model_crud():
    """Verify ClassRoom model CRUD, level, and Major foreign key relation."""
    major = Major.objects.create(code='TBSM', name='Teknik Sepeda Motor')
    classroom = ClassRoom.objects.create(
        name='X TBSM 1',
        level=10,
        major=major
    )
    assert str(classroom) == 'X TBSM 1'
    assert classroom.major == major
    assert classroom.level == 10
    assert major.classes.count() == 1


@pytest.mark.django_db
def test_teacher_model_crud():
    """Verify Teacher model CRUD, CustomUser link, and role level."""
    user = User.objects.create_user(
        username='budi_guru',
        email='budi@smkn1rongga.sch.id',
        full_name='Budi Santoso, S.Pd',
        role=Role.GURU
    )
    teacher = Teacher.objects.create(
        user=user,
        nip='198501012010011001',
        full_name='Budi Santoso, S.Pd',
        email='budi@smkn1rongga.sch.id',
        role_level='guru'
    )
    assert str(teacher) == 'Budi Santoso, S.Pd (Guru)'
    assert teacher.user == user
    assert user.teacher_profile == teacher


@pytest.mark.django_db
def test_student_model_crud():
    """Verify Student model CRUD, CustomUser link, Class & Major FKs."""
    major = Major.objects.create(code='RPL', name='Rekayasa Perangkat Lunak')
    classroom = ClassRoom.objects.create(name='XII RPL 1', level=12, major=major)
    
    user = User.objects.create_user(
        username='22231001',
        nis='22231001',
        full_name='Ahmad Fauzan',
        role=Role.SISWA
    )
    student = Student.objects.create(
        user=user,
        nis='22231001',
        full_name='Ahmad Fauzan',
        major=major,
        class_room=classroom,
        status='aktif',
        email='22231001@student.smkn1rongga.sch.id',
        password_plain='jingga22231001'
    )
    assert str(student) == '22231001 - Ahmad Fauzan'
    assert student.major == major
    assert student.class_room == classroom
    assert student.user == user
    assert student.password_plain == 'jingga22231001'


@pytest.mark.django_db
def test_subject_and_teacher_assignment_crud():
    """Verify Subject model and TeacherAssignment auto-mapping."""
    teacher = Teacher.objects.create(
        full_name='Siti Rahma, M.Kom',
        email='siti@smkn1rongga.sch.id',
        role_level='guru'
    )
    classroom = ClassRoom.objects.create(name='XI RPL 2', level=11)
    subject = Subject.objects.create(code='PWPB', name='Pemrograman Web dan Perangkat Bergerak')

    # Assignment with Subject instance
    assignment = TeacherAssignment.objects.create(
        teacher=teacher,
        subject=subject,
        class_room=classroom
    )
    assert assignment.subject_name == 'Pemrograman Web dan Perangkat Bergerak'
    assert 'Siti Rahma' in str(assignment)

    # Assignment with subject_name auto creating Subject
    assignment2 = TeacherAssignment.objects.create(
        teacher=teacher,
        subject_name='Basis Data Terapan',
        class_room=classroom
    )
    assert assignment2.subject is not None
    assert assignment2.subject.name == 'Basis Data Terapan'


# ==============================================================================
# 2. Search & Filtering Queries Tests
# ==============================================================================

@pytest.mark.django_db
def test_student_search_and_filter_queries():
    """Verify searching and filtering students by NIS, Name, Class, and Major."""
    major_rpl = Major.objects.create(code='RPL', name='Rekayasa Perangkat Lunak')
    major_tsm = Major.objects.create(code='TBSM', name='Teknik Sepeda Motor')
    
    cls_rpl = ClassRoom.objects.create(name='X RPL 1', level=10, major=major_rpl)
    cls_tsm = ClassRoom.objects.create(name='X TBSM 1', level=10, major=major_tsm)

    Student.objects.create(nis='1001', full_name='Aditya Pratama', class_room=cls_rpl, major=major_rpl)
    Student.objects.create(nis='1002', full_name='Bambang Pamungkas', class_room=cls_rpl, major=major_rpl)
    Student.objects.create(nis='1003', full_name='Citra Lestari', class_room=cls_tsm, major=major_tsm)

    # Search by NIS
    res_nis = Student.objects.filter(nis__icontains='1001')
    assert res_nis.count() == 1
    assert res_nis.first().full_name == 'Aditya Pratama'

    # Search by Name
    res_name = Student.objects.filter(full_name__icontains='bambang')
    assert res_name.count() == 1

    # Filter by Class
    res_class = Student.objects.filter(class_room=cls_rpl)
    assert res_class.count() == 2

    # Filter by Major
    res_major = Student.objects.filter(major=major_tsm)
    assert res_major.count() == 1
    assert res_major.first().nis == '1003'


# ==============================================================================
# 3. OpenAPI Synchronization Service Tests (Mocked)
# ==============================================================================

@pytest.mark.django_db
def test_openapi_sync_majors_service():
    """Verify sync_majors extracts unique majors from raw students data."""
    service = MasterDataSyncService(base_url='https://mock.data.smkn1rongga.sch.id')

    mock_students = [
        {'nis': '101', 'name': 'Siswa 1', 'major': 'Rekayasa Perangkat Lunak'},
        {'nis': '102', 'name': 'Siswa 2', 'major': 'Teknik Sepeda Motor'},
        {'nis': '103', 'name': 'Siswa 3', 'major': 'Teknik Kendaraan Ringan Otomotif'},
        {'nis': '104', 'name': 'Siswa 4', 'major': 'Agribisnis Tanaman Pangan dan Hortikultura'},
    ]

    res = service.sync_majors(raw_students=mock_students)
    assert res['success'] is True
    assert res['count'] == 4
    assert Major.objects.filter(code='RPL').exists()
    assert Major.objects.filter(code='TBSM').exists()
    assert Major.objects.filter(code='TKRO').exists()
    assert Major.objects.filter(code='ATPH').exists()


@pytest.mark.django_db
def test_openapi_sync_classes_service():
    """Verify sync_classes correctly extracts classes, assigns level and links majors."""
    service = MasterDataSyncService()

    mock_students = [
        {'nis': '101', 'class_name': 'X RPL 1'},
        {'nis': '102', 'class_name': 'XI TBSM 2'},
        {'nis': '103', 'class_name': 'XII TKRO 1'},
        {'nis': '104', 'class_name': 'X ATPH'},
    ]

    res = service.sync_classes(raw_students=mock_students)
    assert res['success'] is True
    assert res['count'] == 4

    c_rpl = ClassRoom.objects.get(name='X RPL 1')
    assert c_rpl.level == 10
    assert c_rpl.major.code == 'RPL'

    c_tbsm = ClassRoom.objects.get(name='XI TBSM 2')
    assert c_tbsm.level == 11
    assert c_tbsm.major.code == 'TBSM'

    c_tkro = ClassRoom.objects.get(name='XII TKRO 1')
    assert c_tkro.level == 12
    assert c_tkro.major.code == 'TKRO'


@pytest.mark.django_db
def test_openapi_sync_teachers_service():
    """Verify sync_teachers creates Teacher records and CustomUser accounts."""
    service = MasterDataSyncService()

    mock_staff = [
        {
            'id': str(uuid.uuid4()),
            'nip': '197508102005011002',
            'full_name': 'Drs. Supriyadi, M.Pd',
            'email': 'supriyadi@smkn1rongga.sch.id',
            'role_level': 'admin'
        },
        {
            'id': str(uuid.uuid4()),
            'nip': '198203152010012005',
            'full_name': 'Dewi Sartika, S.Kom',
            'email': 'dewi@smkn1rongga.sch.id',
            'role_level': 'guru'
        }
    ]

    res = service.sync_teachers(raw_staff=mock_staff)
    assert res['success'] is True
    assert res['count'] == 2

    # Check Teacher & CustomUser creation
    t_admin = Teacher.objects.get(email='supriyadi@smkn1rongga.sch.id')
    assert t_admin.role_level == 'admin'
    assert t_admin.user is not None
    assert t_admin.user.role == Role.ADMIN
    assert t_admin.user.is_admin is True

    t_guru = Teacher.objects.get(email='dewi@smkn1rongga.sch.id')
    assert t_guru.role_level == 'guru'
    assert t_guru.user.role == Role.GURU


@pytest.mark.django_db
def test_openapi_sync_students_service():
    """Verify sync_students creates Student records, links Class & Major, and generates default credentials."""
    service = MasterDataSyncService()

    mock_students = [
        {
            'nis': '23241001',
            'full_name': 'Bayu Permana',
            'class_name': 'X RPL 1',
            'email': '23241001@student.smkn1rongga.sch.id',
            'status': 'aktif'
        },
        {
            'nis': '23241002',
            'full_name': 'Cindy Claudia',
            'class_name': 'XI TBSM 1',
            'email': '23241002@student.smkn1rongga.sch.id',
            'status': 'aktif'
        }
    ]

    res = service.sync_students(raw_students=mock_students)
    assert res['success'] is True
    assert res['count'] == 2

    s1 = Student.objects.get(nis='23241001')
    assert s1.full_name == 'Bayu Permana'
    assert s1.class_room.name == 'X RPL 1'
    assert s1.major.code == 'RPL'
    assert s1.user is not None
    assert s1.user.username == '23241001'
    assert s1.user.role == Role.SISWA
    assert s1.user.check_password('jingga23241001') is True


@pytest.mark.django_db
@patch.object(MasterDataSyncService, 'fetch_master_api')
def test_openapi_sync_all_service(mock_fetch):
    """Verify sync_all orchestrates majors, classes, teachers, and students."""
    mock_students = [
        {'nis': '3001', 'name': 'Siswa A', 'class_name': 'XII RPL 2', 'major': 'Rekayasa Perangkat Lunak'}
    ]
    mock_staff = [
        {'nip': '12345', 'full_name': 'Guru A', 'email': 'guru_a@smkn1rongga.sch.id', 'role_level': 'guru'}
    ]

    def mock_api_call(endpoint, **kwargs):
        if 'staff' in endpoint or 'users' in endpoint:
            return mock_staff
        return mock_students

    mock_fetch.side_effect = mock_api_call

    service = MasterDataSyncService()
    res = service.sync_all()

    assert res['success'] is True
    assert res['majors_count'] >= 1
    assert res['classes_count'] >= 1
    assert res['teachers_count'] == 1
    assert res['students_count'] == 1


# ==============================================================================
# 4. Views, HTMX, and Export Endpoints Tests
# ==============================================================================

@pytest.mark.django_db
def test_major_list_view(client):
    """Verify MajorListView renders M3 cards for Admin user."""
    admin_user = User.objects.create_superuser(username='admin', email='admin@test.com', password='password123')
    client.force_login(admin_user)

    Major.objects.create(code='RPL', name='Rekayasa Perangkat Lunak')
    Major.objects.create(code='TBSM', name='Teknik Sepeda Motor')

    response = client.get(reverse('master_data:majors'))
    assert response.status_code == 200
    assert 'Data Konsentrasi Keahlian / Jurusan' in response.content.decode()
    assert 'Rekayasa Perangkat Lunak' in response.content.decode()
    assert 'Teknik Sepeda Motor' in response.content.decode()


@pytest.mark.django_db
def test_class_room_list_view_and_filtering(client):
    """Verify ClassRoomListView filtering per major and search."""
    admin_user = User.objects.create_superuser(username='admin', email='admin@test.com', password='password123')
    client.force_login(admin_user)

    m1 = Major.objects.create(code='RPL', name='Rekayasa Perangkat Lunak')
    m2 = Major.objects.create(code='TBSM', name='Teknik Sepeda Motor')
    ClassRoom.objects.create(name='X RPL 1', level=10, major=m1)
    ClassRoom.objects.create(name='XI TBSM 1', level=11, major=m2)

    # Search filter
    response = client.get(reverse('master_data:classes') + '?q=RPL')
    assert response.status_code == 200
    content = response.content.decode()
    assert 'X RPL 1' in content
    assert 'XI TBSM 1' not in content


@pytest.mark.django_db
def test_teacher_list_view_and_export(client):
    """Verify TeacherListView and Excel export."""
    admin_user = User.objects.create_superuser(username='admin', email='admin@test.com', password='password123')
    client.force_login(admin_user)

    Teacher.objects.create(full_name='Budi Santoso', email='budi@smkn1rongga.sch.id', role_level='guru')

    # View
    response = client.get(reverse('master_data:teachers'))
    assert response.status_code == 200
    assert 'Budi Santoso' in response.content.decode()

    # Excel export
    export_res = client.get(reverse('master_data:teachers_export'))
    assert export_res.status_code == 200
    assert export_res['Content-Type'] == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    
    # Check valid openpyxl workbook
    wb = openpyxl.load_workbook(io.BytesIO(export_res.content))
    assert 'Data Guru CBT' in wb.sheetnames


@pytest.mark.django_db
def test_student_list_view_and_htmx_partial(client):
    """Verify StudentListView full page and HTMX partial rendering."""
    admin_user = User.objects.create_superuser(username='admin', email='admin@test.com', password='password123')
    client.force_login(admin_user)

    m = Major.objects.create(code='RPL', name='Rekayasa Perangkat Lunak')
    c = ClassRoom.objects.create(name='X RPL 1', level=10, major=m)
    Student.objects.create(nis='5501', full_name='Dedi Kurniawan', class_room=c, major=m)

    # Full page request
    res_full = client.get(reverse('master_data:students'))
    assert res_full.status_code == 200
    assert 'Data Siswa Peserta CBT' in res_full.content.decode()
    assert 'Dedi Kurniawan' in res_full.content.decode()

    # HTMX partial swap request
    res_htmx = client.get(
        reverse('master_data:students') + '?q=Dedi',
        HTTP_HX_REQUEST='true'
    )
    assert res_htmx.status_code == 200
    htmx_content = res_htmx.content.decode()
    assert 'id="student-table-container"' in htmx_content
    assert 'Dedi Kurniawan' in htmx_content


@pytest.mark.django_db
def test_student_export_excel(client):
    """Verify Student export to Excel."""
    admin_user = User.objects.create_superuser(username='admin', email='admin@test.com', password='password123')
    client.force_login(admin_user)

    Student.objects.create(nis='7701', full_name='Eka Putri', status='aktif')

    res = client.get(reverse('master_data:students_export'))
    assert res.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    assert 'Data Siswa CBT' in wb.sheetnames


@pytest.mark.django_db
def test_teacher_assignment_crud_views(client):
    """Verify TeacherAssignmentListView POST create, update, and delete actions."""
    admin_user = User.objects.create_superuser(username='admin', email='admin@test.com', password='password123')
    client.force_login(admin_user)

    teacher = Teacher.objects.create(full_name='Pak Hendra', email='hendra@test.com')
    subject = Subject.objects.create(name='Matematika Terapan')
    c1 = ClassRoom.objects.create(name='X RPL 1', level=10)
    c2 = ClassRoom.objects.create(name='X RPL 2', level=10)

    # 1. Create assignments
    post_data = {
        'action': 'create',
        'teacher_id': str(teacher.id),
        'subject_id': str(subject.id),
        'class_ids': [str(c1.id), str(c2.id)]
    }
    create_res = client.post(reverse('master_data:assignments'), data=post_data, follow=True)
    assert create_res.status_code == 200
    assert TeacherAssignment.objects.filter(teacher=teacher).count() == 2

    # 2. Update assignment
    assignment = TeacherAssignment.objects.filter(teacher=teacher, class_room=c1).first()
    update_data = {
        'action': 'update',
        'assignment_id': str(assignment.id),
        'teacher_id': str(teacher.id),
        'subject_id': str(subject.id),
        'class_id': str(c2.id)
    }
    update_res = client.post(reverse('master_data:assignments'), data=update_data, follow=True)
    assert update_res.status_code == 200
    assignment.refresh_from_db()
    assert assignment.class_room == c2

    # 3. Delete assignment
    delete_data = {
        'action': 'delete',
        'assignment_id': str(assignment.id)
    }
    delete_res = client.post(reverse('master_data:assignments'), data=delete_data, follow=True)
    assert delete_res.status_code == 200
    assert not TeacherAssignment.objects.filter(id=assignment.id).exists()


@pytest.mark.django_db
@patch.object(MasterDataSyncService, 'sync_all')
def test_sync_master_data_view_htmx(mock_sync, client):
    """Verify SyncMasterDataView handles HTMX POST and renders sync_status partial."""
    admin_user = User.objects.create_superuser(username='admin', email='admin@test.com', password='password123')
    client.force_login(admin_user)

    mock_sync.return_value = {
        'success': True,
        'majors_count': 4,
        'classes_count': 22,
        'teachers_count': 46,
        'students_count': 722,
        'message': 'Seluruh Data Master telah diperbarui.'
    }

    response = client.post(
        reverse('master_data:sync'),
        data={'type': 'all'},
        HTTP_HX_REQUEST='true'
    )
    assert response.status_code == 200
    content = response.content.decode()
    assert 'Sinkronisasi Berhasil' in content
    assert '722 Siswa' in content


# ==============================================================================
# 5. RBAC Security Checks
# ==============================================================================

@pytest.mark.django_db
def test_rbac_security_checks(client):
    """Ensure students cannot access master data and are redirected properly."""
    student_user = User.objects.create_user(
        username='student_user',
        password='password123',
        role=Role.SISWA
    )
    client.force_login(student_user)

    # Siswa accessing master data should be redirected to /student/dashboard/
    for url in [
        reverse('master_data:majors'),
        reverse('master_data:classes'),
        reverse('master_data:teachers'),
        reverse('master_data:students'),
        reverse('master_data:assignments'),
    ]:
        res = client.get(url)
        assert res.status_code == 302
        assert '/student/dashboard/' in res['Location']

    # Unauthenticated user should be redirected to login
    client.logout()
    res_anon = client.get(reverse('master_data:majors'))
    assert res_anon.status_code == 302
    assert reverse('accounts:login') in res_anon['Location']


@pytest.mark.django_db
def test_penugasan_guru_named_url_and_template_rendering(client):
    """Verify /penugasan-guru/ route exists, renders duplicate UI, and handles POST actions."""
    admin_user = User.objects.create_superuser(username='superadmin', email='superadmin@test.com', password='password123')
    client.force_login(admin_user)

    teacher = Teacher.objects.create(full_name='Ibu Siti, M.Pd', email='siti@test.com')
    subject = Subject.objects.create(name='Bahasa Indonesia')
    c1 = ClassRoom.objects.create(name='X RPL 1', level=10)

    # Access GET /penugasan-guru/
    res = client.get(reverse('penugasan_guru'))
    assert res.status_code == 200
    content = res.content.decode()
    assert 'Penugasan Pengampu' in content
    assert 'Atur guru yang mengampu mata pelajaran di tiap kelas' in content
    assert 'Daftar Pengampu' in content
    assert 'Pilih Guru' in content
    assert 'Pilih Mata Pelajaran' in content
    assert 'Pilih Kelas' in content
    assert 'Simpan Penugasan' in content

    # Test POST create via /penugasan-guru/
    post_res = client.post(reverse('penugasan_guru'), data={
        'action': 'create',
        'teacher_id': str(teacher.id),
        'subject_id': str(subject.id),
        'class_ids': [str(c1.id)]
    }, follow=True)
    assert post_res.status_code == 200
    assert TeacherAssignment.objects.filter(teacher=teacher, subject=subject, class_room=c1).exists()


@pytest.mark.django_db
def test_subject_crud_views_and_template(client):
    """Verify /master-mapel/ route exists, renders M3 cards, and handles add, edit, and delete."""
    admin_user = User.objects.create_superuser(username='mapel_admin', email='mapel_admin@test.com', password='password123')
    client.force_login(admin_user)

    s1 = Subject.objects.create(name='Matematika Wajib')
    s2 = Subject.objects.create(name='Bahasa Inggris')

    # GET /master-mapel/
    res = client.get(reverse('master_mapel'))
    assert res.status_code == 200
    content = res.content.decode()
    assert 'Master Mapel' in content
    assert 'Total: 2 Mata Pelajaran Resmi' in content
    assert 'Matematika Wajib' in content
    assert 'Bahasa Inggris' in content
    assert 'Template' in content
    assert 'Import' in content
    assert 'Tambah' in content

    # POST create
    res_create = client.post(reverse('master_mapel'), data={
        'action': 'create',
        'name': 'Fisika Terapan'
    }, follow=True)
    assert res_create.status_code == 200
    assert Subject.objects.filter(name='Fisika Terapan').exists()

    # POST update
    new_sub = Subject.objects.get(name='Fisika Terapan')
    res_update = client.post(reverse('master_mapel'), data={
        'action': 'update',
        'subject_id': str(new_sub.id),
        'name': 'Fisika Dasar dan Terapan'
    }, follow=True)
    assert res_update.status_code == 200
    new_sub.refresh_from_db()
    assert new_sub.name == 'Fisika Dasar dan Terapan'

    # POST delete
    res_del = client.post(reverse('master_mapel'), data={
        'action': 'delete',
        'subject_id': str(new_sub.id)
    }, follow=True)
    assert res_del.status_code == 200
    assert not Subject.objects.filter(id=new_sub.id).exists()


@pytest.mark.django_db
def test_subject_template_and_import_excel(client):
    """Verify Excel template download and batch Excel import for Master Mapel."""
    admin_user = User.objects.create_superuser(username='excel_admin', email='excel_admin@test.com', password='password123')
    client.force_login(admin_user)

    # 1. Test template download
    res_tpl = client.get(reverse('master_mapel_template'))
    assert res_tpl.status_code == 200
    assert res_tpl['Content-Type'] == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    assert 'Template_Import_Mapel_Jingga.xlsx' in res_tpl['Content-Disposition']

    # 2. Test import excel
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Template Mapel"
    ws.append(["Nama_Mapel"])
    ws.append(["Kimia Industri"])
    ws.append(["Biologi Terapan"])

    excel_buffer = io.BytesIO()
    wb.save(excel_buffer)
    excel_buffer.seek(0)
    excel_buffer.name = "test_import_mapel.xlsx"

    res_imp = client.post(reverse('master_mapel_import'), data={
        'excel_file': excel_buffer
    }, follow=True)
    assert res_imp.status_code == 200
    assert Subject.objects.filter(name='Kimia Industri').exists()
    assert Subject.objects.filter(name='Biologi Terapan').exists()


@pytest.mark.django_db
def test_student_import_template_and_process(client):
    """Verify Excel template download and batch Excel import for Students."""
    admin_user = User.objects.create_superuser(username='student_admin', email='student_admin@test.com', password='password123')
    client.force_login(admin_user)

    major = Major.objects.create(code='RPL', name='Rekayasa Perangkat Lunak')
    classroom = ClassRoom.objects.create(name='X RPL 1', level=10, major=major)

    # 1. Download template
    res_tpl = client.get(reverse('import_siswa_template'))
    assert res_tpl.status_code == 200
    assert 'Template_SIAKAD_Jingga.xlsx' in res_tpl['Content-Disposition']

    # 2. GET import page
    res_page = client.get(reverse('import_siswa'))
    assert res_page.status_code == 200
    assert 'Import Siswa' in res_page.content.decode('utf-8')

    # 3. POST Excel file
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Template_Siswa"
    ws.append(["nis", "nama", "nama_kelas"])
    ws.append(["999001", "SISWA TEST IMPORT 1", "X RPL 1"])
    ws.append(["999002", "SISWA TEST IMPORT 2", "10 RPL 1"])  # test alias normalization

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    buf.name = "siswa_test.xlsx"

    res_post = client.post(reverse('import_siswa'), data={
        'excel_file': buf
    }, follow=True)
    assert res_post.status_code == 200

    s1 = Student.objects.filter(nis="999001").first()
    assert s1 is not None
    assert s1.full_name == "SISWA TEST IMPORT 1"
    assert s1.class_room == classroom
    assert s1.user is not None
    assert s1.user.check_password("jingga999001")

    s2 = Student.objects.filter(nis="999002").first()
    assert s2 is not None
    assert s2.full_name == "SISWA TEST IMPORT 2"
    assert s2.class_room == classroom


@pytest.mark.django_db
def test_exam_participants_alias_route(client):
    """Verify /exam-participants/<schedule_id>/ route resolves properly."""
    admin_user = User.objects.create_superuser(
        username='exam_admin',
        email='exam_admin@test.com',
        password='password123',
        role=Role.ADMIN
    )
    client.force_login(admin_user)

    from apps.schedules.models import Exam, Schedule
    from django.utils import timezone
    from datetime import timedelta

    major = Major.objects.create(code='TBSM', name='Teknik Sepeda Motor')
    classroom = ClassRoom.objects.create(name='XI TBSM 1', level=11, major=major)
    exam = Exam.objects.create(title="Ujian Peserta Test", duration=60)
    now = timezone.now()
    schedule = Schedule.objects.create(
        exam=exam,
        class_room=classroom,
        start_time=now,
        end_time=now + timedelta(hours=2),
        token="ABCDEF"
    )

    url = f"/exam-participants/{schedule.id}/"
    res = client.get(url)
    assert res.status_code == 200
    assert "Session Monitoring" in res.content.decode('utf-8')



