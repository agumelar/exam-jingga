"""Tests for Reports, Logistics, Settings & Live Monitoring Module (Task 8)."""
import io
import uuid
import pytest
from datetime import timedelta
from openpyxl import load_workbook
from django.contrib.auth import get_user_model
from django.urls import reverse
from django.utils import timezone

from apps.accounts.models import Role
from apps.core.models import SchoolSetting
from apps.master_data.models import Subject, Teacher, ClassRoom, Major, Student, TeacherAssignment
from apps.questions.models import Question
from apps.schedules.models import Exam, ExamQuestion, Schedule
from apps.exam_engine.models import ExamSession, StudentAnswer
from apps.reports.models import StudentLogistic
from apps.reports.utils import (
    calculate_distractor_analysis,
    auto_distribute_logistics,
    generate_exam_results_excel,
    generate_logistics_excel,
)

User = get_user_model()


@pytest.fixture
def test_setup():
    """Fixture to set up database objects for reports & monitoring tests."""
    major = Major.objects.create(code='RPL', name='Rekayasa Perangkat Lunak')
    classroom_10 = ClassRoom.objects.create(name='X RPL 1', level=10, major=major)
    classroom_11 = ClassRoom.objects.create(name='XI RPL 1', level=11, major=major)
    classroom_12 = ClassRoom.objects.create(name='XII RPL 1', level=12, major=major)

    subject = Subject.objects.create(name='Pemrograman Berorientasi Objek')
    
    # Admin User
    admin_user = User.objects.create_user(
        username='admin_test',
        password='password123',
        role=Role.ADMIN,
        full_name='Administrator CBT'
    )
    admin_teacher = Teacher.objects.create(
        user=admin_user,
        full_name='Administrator CBT',
        role_level='admin'
    )

    # Teacher User
    teacher_user = User.objects.create_user(
        username='guru_test',
        password='password123',
        role=Role.GURU,
        full_name='Pak Guru RPL'
    )
    teacher = Teacher.objects.create(
        user=teacher_user,
        full_name='Pak Guru RPL',
        role_level='guru'
    )

    # Assignment
    TeacherAssignment.objects.create(
        teacher=teacher,
        subject=subject,
        class_room=classroom_12
    )

    # Students
    student1_user = User.objects.create_user(
        username='siswa1',
        password='password123',
        role=Role.SISWA,
        full_name='Ahmad Siswa'
    )
    student1 = Student.objects.create(
        user=student1_user,
        nis='10001',
        full_name='Ahmad Siswa',
        class_room=classroom_12,
        major=major,
        password_plain='pass123'
    )

    student2_user = User.objects.create_user(
        username='siswa2',
        password='password123',
        role=Role.SISWA,
        full_name='Budi Siswa'
    )
    student2 = Student.objects.create(
        user=student2_user,
        nis='10002',
        full_name='Budi Siswa',
        class_room=classroom_12,
        major=major,
        password_plain='pass456'
    )

    student3_user = User.objects.create_user(
        username='siswa3',
        password='password123',
        role=Role.SISWA,
        full_name='Citra Siswa'
    )
    student3 = Student.objects.create(
        user=student3_user,
        nis='10003',
        full_name='Citra Siswa',
        class_room=classroom_11,
        major=major,
        password_plain='pass789'
    )

    # Questions (Bank Soal)
    q1 = Question.objects.create(
        subject=subject,
        created_by=teacher,
        level=12,
        question_text='Apa itu enkapsulasi dalam OOP?',
        option_a='Pembungkusan data dan method',
        option_b='Pewarisan sifat class',
        option_c='Banyak bentuk class',
        option_d='Penyembunyian hardware',
        option_e='Pengacakan data array',
        correct_answer='A'
    )
    q2 = Question.objects.create(
        subject=subject,
        created_by=teacher,
        level=12,
        question_text='Keyword untuk inheritance di Java?',
        option_a='implements',
        option_b='extends',
        option_c='inherits',
        option_d='instanceof',
        option_e='super',
        correct_answer='B'
    )

    # Exam
    exam = Exam.objects.create(
        teacher=teacher,
        subject=subject,
        title='UAS PBO Kelas 12',
        exam_type='PAS',
        duration=60,
        target_question_count=2,
        level=12,
        status='ready',
        token='PBOPAS'
    )
    ExamQuestion.objects.create(exam=exam, question=q1, order_number=1)
    ExamQuestion.objects.create(exam=exam, question=q2, order_number=2)

    # Schedule
    schedule = Schedule.objects.create(
        exam=exam,
        class_room=classroom_12,
        teacher=teacher,
        start_time=timezone.now() - timedelta(minutes=30),
        end_time=timezone.now() + timedelta(hours=2),
        token='PBOPAS',
        session_no=1,
        room_name='Lab RPL 1',
        status='active'
    )

    # Logistics
    log1 = StudentLogistic.objects.create(
        student=student1,
        room_name='Lab RPL 1',
        session_name='SESI 1',
        exam_period='PAS Ganjil 2025/2026'
    )
    log2 = StudentLogistic.objects.create(
        student=student2,
        room_name='Lab RPL 1',
        session_name='SESI 1',
        exam_period='PAS Ganjil 2025/2026'
    )

    return {
        'major': major,
        'classroom_10': classroom_10,
        'classroom_11': classroom_11,
        'classroom_12': classroom_12,
        'subject': subject,
        'admin_user': admin_user,
        'teacher_user': teacher_user,
        'student1': student1,
        'student2': student2,
        'student3': student3,
        'q1': q1,
        'q2': q2,
        'exam': exam,
        'schedule': schedule,
        'log1': log1,
        'log2': log2,
    }


# ==============================================================================
# 1. Models & Logistics Tests
# ==============================================================================

@pytest.mark.django_db
def test_student_logistic_model_creation(test_setup):
    """Verify StudentLogistic model creation, properties, and representation."""
    log = test_setup['log1']
    assert isinstance(log.id, uuid.UUID)
    assert log.room_name == 'Lab RPL 1'
    assert log.session_name == 'SESI 1'
    assert log.student == test_setup['student1']
    assert 'Ahmad Siswa' in str(log)


@pytest.mark.django_db
def test_auto_distribute_logistics(test_setup):
    """Verify auto-distribution of students across rooms and sessions."""
    allocated_count = auto_distribute_logistics(
        levels=[10, 11, 12],
        room_count=2,
        capacity=2,
        sessions_count=2,
        exam_period='PAS Ganjil 2025/2026'
    )
    assert allocated_count == 3  # student1, student2, student3
    assert StudentLogistic.objects.count() == 3
    assert StudentLogistic.objects.filter(room_name='RUANG 01').exists()


# ==============================================================================
# 2. Session Monitoring Views & HTMX Actions
# ==============================================================================

@pytest.mark.django_db
def test_session_monitoring_view_and_table_htmx(client, test_setup):
    """Verify SessionMonitoringView and partial table rendering."""
    client.force_login(test_setup['admin_user'])

    # Create active and locked sessions
    s1 = ExamSession.objects.create(
        student=test_setup['student1'],
        schedule=test_setup['schedule'],
        status='active',
        violation_count=2
    )
    s2 = ExamSession.objects.create(
        student=test_setup['student2'],
        schedule=test_setup['schedule'],
        status='locked',
        violation_count=3
    )

    # Full view GET
    response = client.get('/session-management/')
    assert response.status_code == 200
    assert 'Session Monitoring' in response.content.decode()
    assert 'Ahmad Siswa' in response.content.decode()
    assert 'Budi Siswa' in response.content.decode()

    # HTMX partial GET
    response_htmx = client.get(
        '/session-management/table/',
        HTTP_HX_REQUEST='true'
    )
    assert response_htmx.status_code == 200
    content = response_htmx.content.decode()
    assert 'Mengerjakan' in content
    assert 'Terkunci' in content


@pytest.mark.django_db
def test_session_unlock_reset_and_force_submit_actions(client, test_setup):
    """Verify unlock, reset, and force submit actions on ExamSession."""
    client.force_login(test_setup['admin_user'])

    session = ExamSession.objects.create(
        student=test_setup['student1'],
        schedule=test_setup['schedule'],
        status='locked',
        violation_count=3
    )

    # 1. Unlock Action
    unlock_url = f'/session-management/unlock/{session.id}/'
    res_unlock = client.post(unlock_url, HTTP_HX_REQUEST='true')
    assert res_unlock.status_code == 200
    session.refresh_from_db()
    assert session.status == 'active'
    assert session.violation_count == 0

    # 2. Force Submit Action
    # Give student answers: Q1 correct, Q2 wrong
    StudentAnswer.objects.create(
        session=session,
        question=test_setup['q1'],
        chosen_answer='A'
    )
    StudentAnswer.objects.create(
        session=session,
        question=test_setup['q2'],
        chosen_answer='C'
    )
    force_url = f'/session-management/force-submit/{session.id}/'
    res_force = client.post(force_url, HTTP_HX_REQUEST='true')
    assert res_force.status_code == 200
    session.refresh_from_db()
    assert session.status == 'finished'
    assert session.score == 50.0  # 1 of 2 questions correct

    # 3. Reset Action
    reset_url = f'/session-management/reset/{session.id}/'
    res_reset = client.post(reset_url, HTTP_HX_REQUEST='true')
    assert res_reset.status_code == 200
    assert not ExamSession.objects.filter(id=session.id).exists()


# ==============================================================================
# 3. Exam Results & Distractor Analysis Tests
# ==============================================================================

@pytest.mark.django_db
def test_distractor_analysis_calculation(test_setup):
    """Verify statistical calculation of choices distribution, difficulty level and correctness."""
    exam = test_setup['exam']
    schedule = test_setup['schedule']

    # Student 1 finishes exam (answers Q1: A (correct), Q2: B (correct)) -> score 100
    sess1 = ExamSession.objects.create(
        student=test_setup['student1'],
        schedule=schedule,
        status='finished',
        score=100.0
    )
    StudentAnswer.objects.create(session=sess1, question=test_setup['q1'], chosen_answer='A', is_correct=True)
    StudentAnswer.objects.create(session=sess1, question=test_setup['q2'], chosen_answer='B', is_correct=True)

    # Student 2 finishes exam (answers Q1: A (correct), Q2: C (wrong)) -> score 50
    sess2 = ExamSession.objects.create(
        student=test_setup['student2'],
        schedule=schedule,
        status='finished',
        score=50.0
    )
    StudentAnswer.objects.create(session=sess2, question=test_setup['q1'], chosen_answer='A', is_correct=True)
    StudentAnswer.objects.create(session=sess2, question=test_setup['q2'], chosen_answer='C', is_correct=False)

    analysis_res = calculate_distractor_analysis(exam)
    assert analysis_res['total_finished'] == 2
    assert analysis_res['questions_count'] == 2

    q1_data = analysis_res['analysis_data'][0]
    assert q1_data['correct'] == 2
    assert q1_data['wrong'] == 0
    assert q1_data['distro']['A'] == 2
    assert q1_data['correct_percentage'] == 100.0
    assert q1_data['difficulty_level'] == 'Mudah'

    q2_data = analysis_res['analysis_data'][1]
    assert q2_data['correct'] == 1
    assert q2_data['wrong'] == 1
    assert q2_data['distro']['B'] == 1
    assert q2_data['distro']['C'] == 1
    assert q2_data['correct_percentage'] == 50.0
    assert q2_data['difficulty_level'] == 'Sedang'


@pytest.mark.django_db
def test_exam_results_view_and_summary_statistics(client, test_setup):
    """Verify ExamResultsView rendering, ranking, passing rate (KKM 75), and filters."""
    client.force_login(test_setup['admin_user'])

    exam = test_setup['exam']
    schedule = test_setup['schedule']

    sess1 = ExamSession.objects.create(
        student=test_setup['student1'],
        schedule=schedule,
        status='finished',
        score=90.0
    )
    sess2 = ExamSession.objects.create(
        student=test_setup['student2'],
        schedule=schedule,
        status='finished',
        score=60.0
    )

    response = client.get(f'/exam-results/{exam.id}/')
    assert response.status_code == 200
    content = response.content.decode()
    assert 'UAS PBO Kelas 12' in content
    assert '75.0' in content or '75' in content  # Average 75
    assert '90' in content  # Highest
    assert '60' in content  # Lowest
    assert 'TUNTAS' in content
    assert 'BELUM TUNTAS' in content


@pytest.mark.django_db
def test_exam_results_excel_export(client, test_setup):
    """Verify Excel workbook generation with openpyxl for exam results."""
    client.force_login(test_setup['admin_user'])

    exam = test_setup['exam']
    schedule = test_setup['schedule']

    ExamSession.objects.create(
        student=test_setup['student1'],
        schedule=schedule,
        status='finished',
        score=85.0
    )

    response = client.get(f'/exam-results/{exam.id}/export/')
    assert response.status_code == 200
    assert response['Content-Type'] == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    assert 'attachment; filename=' in response['Content-Disposition']

    # Load workbook and verify sheets
    wb = load_workbook(io.BytesIO(response.content))
    assert 'Rekap Nilai' in wb.sheetnames
    assert 'Analisis Pengecoh' in wb.sheetnames

    ws_rekap = wb['Rekap Nilai']
    assert ws_rekap.cell(row=4, column=3).value == 'Nama Lengkap Siswa'


# ==============================================================================
# 4. Printable Templates & Logistics Tests
# ==============================================================================

@pytest.mark.django_db
def test_exam_cards_view_rendering(client, test_setup):
    """Verify ExamCardsView rendering with school settings and student credentials."""
    client.force_login(test_setup['admin_user'])

    response = client.get('/exam-cards/')
    assert response.status_code == 200
    content = response.content.decode()
    assert 'Kartu Peserta Ujian' in content
    assert 'Ahmad Siswa' in content
    assert 'pass123' in content
    assert 'Lab RPL 1' in content


@pytest.mark.django_db
def test_attendance_list_view_rendering(client, test_setup):
    """Verify AttendanceListView rendering with letterhead and student groups."""
    client.force_login(test_setup['admin_user'])

    response = client.get('/attendance-list/')
    assert response.status_code == 200
    content = response.content.decode()
    assert 'Daftar Hadir' in content
    assert 'Lab RPL 1' in content
    assert 'SESI 1' in content
    assert 'Ahmad Siswa' in content


@pytest.mark.django_db
def test_logistics_view_and_export(client, test_setup):
    """Verify LogisticsView and Excel export."""
    client.force_login(test_setup['admin_user'])

    # View GET
    response = client.get('/logistics/')
    assert response.status_code == 200
    assert 'Logistik &amp; Sesi Ruangan' in response.content.decode() or 'Logistik & Sesi Ruangan' in response.content.decode()

    # Export GET
    res_export = client.get('/reports/logistics/export/')
    assert res_export.status_code == 200
    assert res_export['Content-Type'] == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'

    wb = load_workbook(io.BytesIO(res_export.content))
    assert 'Logistik Ruang & Sesi' in wb.sheetnames


# ==============================================================================
# 5. Settings Management Tests
# ==============================================================================

@pytest.mark.django_db
def test_settings_view_get_and_post_update(client, test_setup):
    """Verify SettingsView form rendering and updating of SchoolSetting model."""
    client.force_login(test_setup['admin_user'])

    # GET
    response_get = client.get('/settings/')
    assert response_get.status_code == 200
    assert 'Konfigurasi Identitas Lembaga' in response_get.content.decode()

    # POST update
    payload = {
        'school_name': 'SMK NEGERI 1 RONGGA JAYA',
        'school_majors_list': 'RPL, TBSM, TKRO, ATPH',
        'school_address': 'Jl. Bojonghaleuang No. 1, Rongga',
        'school_phone': '(022) 1234567',
        'school_postal_code': '40565',
        'school_website': 'https://smkn1rongga.sch.id',
        'school_email': 'info@smkn1rongga.sch.id',
        'header_1': 'PEMERINTAH DAERAH PROVINSI JAWA BARAT',
        'header_2': 'DINAS PENDIDIKAN',
        'header_3': 'CABANG DINAS PENDIDIKAN WILAYAH VI',
        'headmaster_name': 'Drs. H. Juandi, M.Pd.',
        'headmaster_nip': '196708151994031005',
        'curriculum_vicedir_name': 'Suhendar Aryadi, S.Kom.',
        'curriculum_vicedir_nip': '198501012010011002',
        'committee_chairman': 'Dedi Suhendar, S.Pd.',
        'exam_city': 'Rongga',
        'exam_date': '15 Maret 2026',
        'academic_year': '2025/2026',
        'semester': 'Genap',
        'exam_name': 'Penilaian Akhir Tahun (PAT)',
    }

    response_post = client.post('/settings/', data=payload)
    assert response_post.status_code in [200, 302]

    setting = SchoolSetting.get_settings()
    assert setting.school_name == 'SMK NEGERI 1 RONGGA JAYA'
    assert setting.headmaster_name == 'Drs. H. Juandi, M.Pd.'
    assert setting.curriculum_vicedir_name == 'Suhendar Aryadi, S.Kom.'
    assert setting.exam_name == 'Penilaian Akhir Tahun (PAT)'
    assert setting.academic_year == '2025/2026'
